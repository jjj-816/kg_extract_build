"""人工复核追加记录与首版正式报告导出。"""

from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from docx import Document


class ReportPublishBlocked(ValueError):
    pass


class ReportReauditRequired(ValueError):
    pass


def append_review_action(report: dict, *, task_id: str, action: str, reviewer: str, explanation: str, value: Any = None) -> dict:
    result = deepcopy(report)
    result.setdefault("human_actions", []).append({
        "task_id": task_id, "action": action, "reviewer": reviewer,
        "explanation": explanation, "value": value,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    for task in result.get("tasks", []):
        if task.get("task_id") == task_id:
            task["human_review_status"] = "resolved"
            break
    return result


def _pending_reviews(report: dict) -> list[str]:
    pending = []
    for task in report.get("tasks", []):
        required = bool(task.get("manual_reviews")) or task.get("result_status") == "manual_review"
        resolved = task.get("human_review_status") == "resolved"
        if required and not resolved:
            pending.append(str(task.get("task_id")))
    return pending


def freeze_report_snapshot(report: dict, *, publisher: str) -> dict:
    pending = _pending_reviews(report)
    if pending:
        raise ReportPublishBlocked("以下任务仍有未处理人工必办项：" + ", ".join(pending))
    snapshot = deepcopy(report)
    snapshot["report_type"] = "audit_formal_v1"
    snapshot["publication"] = {"publisher": publisher, "published_at": datetime.now(timezone.utc).isoformat()}
    encoded = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    snapshot["snapshot_sha256"] = hashlib.sha256(encoded).hexdigest()
    return snapshot


def export_report_snapshot(snapshot: dict, output_dir: str | Path) -> tuple[Path, Path]:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    stem = str(snapshot.get("document", {}).get("document_id", "audit-report"))
    json_path = directory / f"{stem}.json"
    docx_path = directory / f"{stem}.docx"
    json_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    document = Document()
    document.add_heading("施工方案正式审核报告", level=1)
    document.add_paragraph(f"报告快照：{snapshot.get('snapshot_sha256', '')}")
    document.add_paragraph(f"发布者：{snapshot.get('publication', {}).get('publisher', '')}")
    for task in snapshot.get("tasks", []):
        document.add_heading(f"{task.get('task_id')} {task.get('task_name', '')}", level=2)
        document.add_paragraph(f"结论：{task.get('result_status') or task.get('execution_status')}")
        for collection in ("issues", "manual_reviews", "offline_items", "advisories"):
            for item in task.get(collection, []):
                document.add_paragraph(str(item.get("summary", "")), style="List Bullet")
    document.save(docx_path)
    return json_path, docx_path


def create_correction_version(previous: dict, *, changes: list[dict], reason: str, source_snapshot: dict) -> dict:
    """Create an append-only human correction version without machine rerun."""
    expected = previous.get("source_snapshot")
    if expected is not None and expected != source_snapshot:
        raise ReportReauditRequired("审核输入、上下文或证据快照已变化，必须新建审核运行")
    result = deepcopy(previous)
    result["report_type"] = "audit_correction_v1"
    result["previous_report_id"] = previous.get("report_id") or previous.get("snapshot_sha256")
    result["version"] = int(previous.get("version", 1)) + 1
    result["change_reason"] = reason
    result["corrections"] = deepcopy(changes)
    result["source_snapshot"] = deepcopy(source_snapshot)
    result["created_at"] = datetime.now(timezone.utc).isoformat()
    for change in changes:
        for task in result.get("tasks", []):
            if task.get("task_id") == change.get("task_id"):
                task.setdefault("human_actions", []).append(change)
                if "result_status" in change:
                    task["result_status"] = change["result_status"]
    encoded = json.dumps(result, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    result["snapshot_sha256"] = hashlib.sha256(encoded).hexdigest()
    return result
