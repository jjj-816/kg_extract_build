"""固定风险作业目录加载及编号到作业类型映射。"""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
import hashlib
from pathlib import Path
import re

from .settings import AUDIT_RISK_CATALOG_PATH


class RiskCatalogError(ValueError):
    """风险作业目录不可用于自动审核。"""


WORK_TYPE_ALIASES = {
    "动火作业": "动火作业", "受限空间作业": "受限空间作业", "临时用电作业": "临时用电作业",
    "动土作业": "动土作业", "高处作业": "高处作业", "吊装作业": "吊装作业", "上位系统类作业": "上位系统类作业",
}
SHEET_WORK_TYPE_HINTS = {"通讯与自控": "上位系统类作业"}
RISK_WORK_CODE_PATTERN = re.compile(r"\b[A-Z]{2,}(?:-[A-Z0-9]+)+\b")


def _normalize(value: object) -> str:
    return "" if value is None else "".join(str(value).split())


def _canonical_work_type(value: object) -> str | None:
    normalized = _normalize(value)
    return next((canonical for alias, canonical in WORK_TYPE_ALIASES.items() if alias in normalized), None)


@dataclass(frozen=True)
class RiskCatalog:
    """由目录编号索引的作业类型集合；类型首先来自“子系统”列。"""

    source_path: Path
    version: str
    work_types_by_code: dict[str, frozenset[str]]

    def work_types_for(self, code: str) -> frozenset[str] | None:
        return self.work_types_by_code.get(code.upper())


@dataclass(frozen=True)
class DetectedWorkCode:
    code: str
    source_text: str
    section: str
    source_locator: str
    work_types: tuple[str, ...]
    catalog_match_status: str

    def as_dict(self) -> dict:
        return {
            "code": self.code, "source_text": self.source_text, "section": self.section,
            "source_locator": self.source_locator, "work_types": list(self.work_types),
            "catalog_match_status": self.catalog_match_status,
        }


def _resolve_catalog_path(path: str | Path | None) -> Path:
    candidate = Path(path or AUDIT_RISK_CATALOG_PATH).expanduser().resolve()
    if candidate.is_file():
        return candidate
    if not candidate.is_dir():
        raise RiskCatalogError(f"风险作业目录文件不存在：{candidate}")
    files = sorted(candidate.glob("*.xlsx"))
    if len(files) != 1:
        raise RiskCatalogError(f"风险作业目录目录应恰好包含一个 .xlsx 文件：{candidate}")
    return files[0]


def _header_row(sheet) -> tuple[int, int, int]:
    for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True), 1):
        normalized = [_normalize(value) for value in row]
        code_column = next((index for index, value in enumerate(normalized) if "风险作业目录序号" in value), None)
        subsystem_column = next((index for index, value in enumerate(normalized) if value == "子系统"), None)
        if code_column is not None and subsystem_column is not None:
            return row_number, code_column, subsystem_column
    raise RiskCatalogError(f"工作表“{sheet.title}”缺少“子系统”或“风险作业目录序号”列")


@lru_cache(maxsize=4)
def _load_catalog(path_text: str) -> RiskCatalog:
    path = Path(path_text)
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise RiskCatalogError(f"无法读取风险作业目录：{path}") from exc
    code_types: dict[str, set[str]] = {}
    try:
        for sheet in workbook.worksheets:
            header_row, code_column, subsystem_column = _header_row(sheet)
            type_header = next(sheet.iter_rows(min_row=header_row + 1, max_row=header_row + 1, values_only=True))
            matrix_columns = {
                index: work_type for index, value in enumerate(type_header)
                if (work_type := _canonical_work_type(value)) is not None
            }
            for row in sheet.iter_rows(min_row=header_row + 3, values_only=True):
                code = _normalize(row[code_column] if code_column < len(row) else None).upper()
                if not code:
                    continue
                work_types = set()
                subsystem_type = _canonical_work_type(row[subsystem_column] if subsystem_column < len(row) else None)
                if subsystem_type:
                    work_types.add(subsystem_type)
                for column, work_type in matrix_columns.items():
                    marker = _normalize(row[column] if column < len(row) else None)
                    if marker and marker not in {"/", "-", "无"}:
                        work_types.add(work_type)
                for hint, work_type in SHEET_WORK_TYPE_HINTS.items():
                    if hint in sheet.title:
                        work_types.add(work_type)
                # 目录按专业拆分工作表，少量编号会跨表重复出现；同一编号
                # 的作业类型取并集，避免因工作表分区导致漏映射。
                code_types.setdefault(code, set()).update(work_types)
    finally:
        workbook.close()
    if not code_types:
        raise RiskCatalogError("风险作业目录未读取到任何编号")
    version = f"{path.name}:{hashlib.sha256(path.read_bytes()).hexdigest()[:12]}"
    return RiskCatalog(path, version, {code: frozenset(types) for code, types in code_types.items()})


def load_risk_catalog(path: str | Path | None = None) -> RiskCatalog:
    """加载项目内目录；环境变量 ``KG_AUDIT_RISK_CATALOG_PATH`` 可覆盖默认路径。"""
    return _load_catalog(str(_resolve_catalog_path(path)))


def detect_work_codes(blocks, catalog: RiskCatalog | None = None) -> tuple[DetectedWorkCode, ...]:
    """从任务证据块提取目录格式编号，并保留每次命中的原文与来源。"""
    catalog = catalog or load_risk_catalog()
    detections = []
    for block in blocks:
        raw_text = block.get("raw_text", "") if isinstance(block, dict) else getattr(block, "raw_text", "")
        section_path = block.get("section_path", ()) if isinstance(block, dict) else getattr(block, "section_path", ())
        source_locator = block.get("source_locator", "") if isinstance(block, dict) else getattr(block, "source_locator", "")
        for code in RISK_WORK_CODE_PATTERN.findall(raw_text.upper()):
            work_types = catalog.work_types_for(code)
            detections.append(DetectedWorkCode(
                code=code,
                source_text=raw_text,
                section=" / ".join(section_path or ()),
                source_locator=source_locator,
                work_types=tuple(sorted(work_types or ())),
                catalog_match_status="matched" if work_types is not None else "unmatched",
            ))
    return tuple(detections)
