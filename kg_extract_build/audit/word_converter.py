"""将传统 .doc 转换为只读解析使用的 .docx 副本。"""

from __future__ import annotations

import hashlib
import os
import platform
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from .settings import AUDIT_CONVERSION_TIMEOUT, AUDIT_DOC_CONVERTER, LIBREOFFICE_PATH


@dataclass(frozen=True)
class WordConversionCapability:
    available: bool
    message: str
    word_com_available: bool = False
    libreoffice_available: bool = False


@dataclass(frozen=True)
class DocxConversionResult:
    path: Path
    content_hash: str
    converter_name: str
    converter_version: str | None
    diagnostics: dict


def doc_conversion_capability() -> WordConversionCapability:
    word_available = False
    if platform.system() == "Windows":
        try:
            import pythoncom  # noqa: F401
            import win32com.client  # noqa: F401
            word_available = True
        except ImportError:
            pass
    libreoffice_available = LIBREOFFICE_PATH.is_file()
    if word_available and libreoffice_available:
        return WordConversionCapability(True, "Word COM 与 LibreOffice 均可用，自动模式支持回退", True, True)
    if word_available:
        return WordConversionCapability(True, "Word COM 可用", True, False)
    if libreoffice_available:
        return WordConversionCapability(True, "LibreOffice 可用", False, True)
    return WordConversionCapability(False, "Word COM 与 LibreOffice 均不可用", False, False)


def _validate_docx(path: Path) -> tuple[str, dict]:
    import zipfile

    if path.suffix.lower() != ".docx" or not path.is_file() or not zipfile.is_zipfile(path):
        raise RuntimeError("转换输出不是有效 DOCX 压缩包")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
    required = {"[Content_Types].xml", "word/document.xml"}
    if not required.issubset(names):
        raise RuntimeError("转换输出缺少必要 DOCX 部件")
    try:
        document = Document(str(path))
    except Exception as exc:
        raise RuntimeError(f"python-docx 无法打开转换输出：{exc}") from exc
    if not document.paragraphs and not document.tables:
        raise RuntimeError("转换输出不包含段落或表格")
    return hashlib.sha256(path.read_bytes()).hexdigest().upper(), {
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "zip_valid": True,
    }


def _word_com_convert(source: Path, output_path: Path) -> tuple[str | None, dict]:
    diagnostics: dict = {"attempted": True}
    started = time.perf_counter()
    word = None
    document = None
    pythoncom = None
    try:
        import pythoncom as pythoncom_module
        import win32com.client

        pythoncom = pythoncom_module
        pythoncom.CoInitialize()
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(
            str(source), ReadOnly=True, AddToRecentFiles=False, OpenAndRepair=True
        )
        document.SaveAs2(str(output_path), FileFormat=16, AddToRecentFiles=False)
        diagnostics["version"] = str(getattr(word, "Version", "")) or None
        return diagnostics["version"], diagnostics
    except Exception as exc:
        diagnostics["error"] = str(exc)
        raise
    finally:
        diagnostics["elapsed_ms"] = int((time.perf_counter() - started) * 1000)
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        if pythoncom is not None:
            pythoncom.CoUninitialize()


def _libreoffice_convert(source: Path, output_dir: Path, timeout: int) -> tuple[str | None, dict]:
    diagnostics: dict = {"attempted": True, "path": str(LIBREOFFICE_PATH)}
    profile_root = Path(tempfile.gettempdir()) / "kg-audit"
    profile_root.mkdir(parents=True, exist_ok=True)
    profile_dir = Path(tempfile.mkdtemp(prefix="lo-profile-", dir=profile_root))
    started = time.perf_counter()
    console_path = LIBREOFFICE_PATH.with_suffix(".com")
    executable = console_path if console_path.is_file() else LIBREOFFICE_PATH
    command = [
        str(executable), "--headless",
        f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
        "--convert-to", "docx", "--outdir", str(output_dir), str(source),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=False, timeout=timeout, check=False)
        stdout = (result.stdout or b"").decode("utf-8", errors="replace")
        stderr = (result.stderr or b"").decode("utf-8", errors="replace")
        diagnostics.update({"executable": str(executable), "stdout": stdout[-4000:], "stderr": stderr[-4000:], "returncode": result.returncode})
        if result.returncode != 0:
            raise RuntimeError(f"LibreOffice 转换失败，返回码 {result.returncode}")
        return None, diagnostics
    except subprocess.TimeoutExpired as exc:
        diagnostics["error"] = f"转换超时：{timeout} 秒"
        diagnostics["stdout"] = (exc.stdout or b"").decode("utf-8", errors="replace")[-4000:]
        diagnostics["stderr"] = (exc.stderr or b"").decode("utf-8", errors="replace")[-4000:]
        raise RuntimeError(diagnostics["error"]) from exc
    finally:
        diagnostics["elapsed_ms"] = int((time.perf_counter() - started) * 1000)
        try:
            shutil.rmtree(profile_dir)
            diagnostics["profile_cleanup"] = "removed"
        except OSError as exc:
            diagnostics["profile_cleanup"] = f"failed: {exc}"


def _wait_for_stable_docx(path: Path, timeout: int, minimum_size: int = 128) -> dict:
    """等待 soffice 可能异步写入的结果完成，避免读取半成品。"""
    started, stable_count, previous = time.monotonic(), 0, None
    checks = 0
    while time.monotonic() - started < timeout:
        checks += 1
        if path.is_file() and path.stat().st_size >= minimum_size:
            stat = path.stat()
            current = (stat.st_size, stat.st_mtime_ns)
            stable_count = stable_count + 1 if current == previous else 0
            previous = current
            if stable_count >= 1:
                try:
                    _validate_docx(path)
                    return {"wait_output_ms": int((time.monotonic() - started) * 1000), "stability_checks": checks}
                except RuntimeError:
                    pass
        time.sleep(0.2)
    raise RuntimeError(f"转换输出在 {timeout} 秒内未生成稳定的有效 DOCX")


def convert_doc_to_docx(
    source_path: str | Path,
    output_dir: str | Path,
    mode: str | None = None,
    timeout: int | None = None,
) -> DocxConversionResult:
    source = Path(source_path).expanduser().resolve()
    destination_dir = Path(output_dir).expanduser().resolve()
    if source.suffix.lower() != ".doc":
        raise ValueError("仅支持将 .doc 转换为 .docx")
    if not source.is_file():
        raise FileNotFoundError(f".doc 文件不存在：{source}")
    selected_mode = (mode or AUDIT_DOC_CONVERTER).lower()
    if selected_mode not in {"auto", "word", "libreoffice"}:
        raise ValueError("KG_AUDIT_DOC_CONVERTER 只能为 auto、word 或 libreoffice")
    capability = doc_conversion_capability()
    if not capability.available:
        raise RuntimeError(capability.message)
    destination_dir.mkdir(parents=True, exist_ok=True)
    output_path = destination_dir / f"{source.stem}.docx"
    diagnostics: dict = {"mode": selected_mode, "word_com": {"attempted": False}, "libreoffice": {"attempted": False}}
    attempts = []
    if selected_mode in {"auto", "word"} and capability.word_com_available:
        attempts.append("word")
    if selected_mode in {"auto", "libreoffice"} and capability.libreoffice_available:
        attempts.append("libreoffice")
    # ``os.replace`` 只支持同一卷；临时转换目录放在当前 document_id 目录下，
    # 既避免跨盘失败，又会在 finally 中清除，不会作为业务文件保留。
    work_dir = Path(tempfile.mkdtemp(prefix=".kg-audit-convert-", dir=destination_dir.parent))
    try:
        for converter in attempts:
            temporary_output = work_dir / f"{source.stem}.docx"
            temporary_output.unlink(missing_ok=True)
            try:
                if converter == "word":
                    diagnostics["word_com"] = {"attempted": True}
                    version, detail = _word_com_convert(source, temporary_output)
                    diagnostics["word_com"] = detail
                else:
                    diagnostics["libreoffice"] = {"attempted": True, "path": str(LIBREOFFICE_PATH)}
                    version, detail = _libreoffice_convert(source, work_dir, timeout or AUDIT_CONVERSION_TIMEOUT)
                    diagnostics["libreoffice"] = detail
                    diagnostics["output_stability"] = _wait_for_stable_docx(temporary_output, timeout or AUDIT_CONVERSION_TIMEOUT)
                content_hash, validation = _validate_docx(temporary_output)
                diagnostics["validation"] = validation
                os.replace(temporary_output, output_path)
                return DocxConversionResult(output_path, content_hash, converter, version, diagnostics)
            except Exception as exc:
                diagnostics[f"{converter}_error"] = str(exc)
                if selected_mode != "auto":
                    break
        raise RuntimeError(".doc 转换失败：" + str(diagnostics))
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)
