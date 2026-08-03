"""按 Word 正文真实顺序解析章节、表格和 DrawingML/VML 图片。"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from .models import AuditDocumentBlock, AuditImage, ParsedAuditDocument, StoredAuditDocument


_CHAPTER_RE = re.compile(r"^第[一二三四五六七八九十百]+章\s+[^。；;]{1,80}$")
_NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){1,3})\s*[^。；;]{1,100}$")
_APPENDIX_RE = re.compile(r"^附录\s*[A-EＡ-Ｅ](?:[：:、\s].*)?$")
_TABLE_APPENDIX_RE = re.compile(r"^\s*(附录\s*[A-EＡ-Ｅ](?:[：:].*)?)", re.MULTILINE)
_STYLE_LEVEL_RE = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)
_NORMALIZE_RE = re.compile(r"[\s\u3000，,。；;：:（）()【】\[\]《》<>‘’'\"、·—-]+")
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_TEMPLATE_CHAPTERS = {
    "编制依据", "工程概况", "施工准备与资源配置计划", "施工安排", "健康、安全、环保管理", "健康安全环保管理", "附录",
}
_SENTENCE_END_RE = re.compile(r"[。；;！？!?]$")


def normalize_for_match(value: str) -> str:
    return _NORMALIZE_RE.sub("", unicodedata.normalize("NFKC", value or "")).lower()


def _iter_block_elements(element):
    """Yield paragraphs/tables from a container, expanding Word content controls."""
    for child in element.iterchildren():
        if isinstance(child, CT_P):
            yield child
        elif isinstance(child, CT_Tbl):
            yield child
        elif child.tag.endswith("}sdt"):
            for nested in child.iterchildren():
                if nested.tag.endswith("}sdtContent"):
                    yield from _iter_block_elements(nested)


def iter_body_blocks(document: DocxDocument) -> Iterable[Paragraph | Table]:
    for child in _iter_block_elements(document.element.body):
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        else:
            yield Table(child, document)


def _iter_container_blocks(container) -> Iterable[Paragraph | Table]:
    element = getattr(container, "element", None)
    if element is None:
        element = getattr(container, "_element", None)
    if isinstance(container, DocxDocument):
        element = container.element.body
    if element is None:
        return
    for child in _iter_block_elements(element):
        if isinstance(child, CT_P):
            yield Paragraph(child, container)
        elif isinstance(child, CT_Tbl):
            yield Table(child, container)


def _is_toc_entry(paragraph: Paragraph, text: str) -> bool:
    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    if style_name.startswith("toc") or style_name.startswith("目录"):
        return True
    xml_text = paragraph._element.xml.lower()
    if " toc " in xml_text or "hyperlink" in xml_text and re.search(r"\t\s*\d+\s*$", text):
        return True
    return bool("\t" in text and re.search(r"\t\s*\d+\s*$", text))


def _heading_level(paragraph: Paragraph, text: str) -> int | None:
    if _is_toc_entry(paragraph, text):
        return None
    if _CHAPTER_RE.match(text) or _APPENDIX_RE.match(text):
        return 1
    numbered = _NUMBERED_HEADING_RE.match(text)
    if numbered:
        return numbered.group(1).count(".") + 1
    normalized = normalize_for_match(text)
    if normalized in {normalize_for_match(value) for value in _TEMPLATE_CHAPTERS}:
        return 1
    # A complete prose sentence is occasionally given a Heading style in
    # source files.  It must not split the evidence region.
    if _SENTENCE_END_RE.search(text) or (len(text) > 80 and not re.match(r"^\d", text)):
        return None
    style_name = getattr(paragraph.style, "name", "") or ""
    style_match = _STYLE_LEVEL_RE.search(style_name)
    if style_match:
        return int(style_match.group(1))
    return None


def _update_section_stack(stack: list[tuple[int, str]], level: int, text: str) -> tuple[str, ...]:
    while stack and stack[-1][0] >= level:
        stack.pop()
    stack.append((level, text))
    return tuple(item[1] for item in stack)


def _relationship_ids(paragraph: Paragraph) -> tuple[str, ...]:
    ids: list[str] = []
    for element in paragraph._element.xpath(".//*[local-name()='blip' or local-name()='imagedata']"):
        relationship_id = element.get(_REL_NS + ("embed" if element.tag.endswith("blip") else "id"))
        if relationship_id and relationship_id not in ids:
            ids.append(relationship_id)
    return tuple(ids)


def _extension_for_content_type(content_type: str) -> str:
    mapping = {
        "image/png": ".png", "image/jpeg": ".jpg", "image/jpg": ".jpg",
        "image/gif": ".gif", "image/bmp": ".bmp", "image/tiff": ".tiff",
        "image/x-emf": ".emf", "image/x-wmf": ".wmf",
    }
    return mapping.get(content_type.lower(), ".bin")


def _extract_images(paragraph: Paragraph, document: StoredAuditDocument, source_part: str, source_locator: str, image_dir: Path, cache: dict[tuple[str, str], AuditImage]) -> tuple[str, ...]:
    image_ids: list[str] = []
    part = paragraph.part
    for relationship_id in _relationship_ids(paragraph):
        cache_key = (str(part.partname), relationship_id)
        image = cache.get(cache_key)
        if image is None:
            try:
                related = part.related_parts[relationship_id]
                content = related.blob
                content_type = related.content_type
            except (KeyError, AttributeError):
                continue
            content_hash = hashlib.sha256(content).hexdigest().upper()
            image_id = f"{document.document_id}-I{len(cache) + 1:04d}"
            extension = _extension_for_content_type(content_type)
            image_dir.mkdir(parents=True, exist_ok=True)
            stored_path = image_dir / f"{image_id}{extension}"
            stored_path.write_bytes(content)
            image = AuditImage(
                image_id=image_id,
                document_id=document.document_id,
                relationship_id=relationship_id,
                source_part=source_part,
                source_locator=source_locator,
                content_type=content_type,
                file_extension=extension,
                content_hash=content_hash,
                stored_path=stored_path,
            )
            cache[cache_key] = image
        image_ids.append(image.image_id)
    return tuple(image_ids)


def _table_payload(table: Table, document: StoredAuditDocument, source_part: str, source_locator: str, image_dir: Path, cache: dict[tuple[str, str], AuditImage]) -> tuple[dict, tuple[str, ...]]:
    rows: list[list[str]] = []
    cells: list[dict] = []
    image_ids: list[str] = []
    seen_vertical = set()
    for row_index, row in enumerate(table.rows, start=1):
        row_values: list[str] = []
        seen_in_row = set()
        for column_index, cell in enumerate(row.cells, start=1):
            tc_identity = cell._tc
            text = "\n".join(p.text.strip() for p in cell.paragraphs).strip()
            merged_continuation = tc_identity in seen_in_row or tc_identity in seen_vertical
            if merged_continuation:
                row_values.append("")
                continue
            seen_in_row.add(tc_identity)
            seen_vertical.add(tc_identity)
            colspan = sum(1 for candidate in row.cells[column_index - 1:] if candidate._tc == tc_identity)
            row_values.append(text)
            cells.append({"row": row_index, "column": column_index - 1, "text": text, "rowspan": 1, "colspan": colspan, "merge_origin": True})
            for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                refs = _extract_images(
                    paragraph, document, source_part,
                    f"{source_locator}/row[{row_index}]/cell[{column_index}]/paragraph[{paragraph_index}]",
                    image_dir, cache,
                )
                for ref in refs:
                    if ref not in image_ids:
                        image_ids.append(ref)
        rows.append(row_values)
    return {"rows": rows, "cells": cells}, tuple(image_ids)


def _parse_container(container, document: StoredAuditDocument, source_part: str, start_ordinal: int, section_stack: list[tuple[int, str]], image_dir: Path, cache: dict[tuple[str, str], AuditImage], include_in_body: bool) -> tuple[list[AuditDocumentBlock], int, int, int]:
    blocks: list[AuditDocumentBlock] = []
    ordinal = start_ordinal
    paragraph_count = 0
    table_count = 0
    for item in _iter_container_blocks(container):
        ordinal += 1
        if isinstance(item, Paragraph):
            paragraph_count += 1
            text = item.text.strip()
            toc_entry = _is_toc_entry(item, text)
            level = _heading_level(item, text) if text else None
            if include_in_body and level is not None:
                section_path = _update_section_stack(section_stack, level, text)
            else:
                section_path = tuple(value for _, value in section_stack)
            refs = _extract_images(item, document, source_part, f"word/{source_part}[{ordinal}]/paragraph", image_dir, cache)
            if not text and not refs:
                continue
            blocks.append(AuditDocumentBlock(
                block_id=f"{document.document_id}-B{ordinal:04d}", document_id=document.document_id,
                ordinal=ordinal, block_type="toc_entry" if toc_entry else ("heading" if level is not None else "paragraph"),
                section_path=section_path, source_locator=f"word/{source_part}[{ordinal}]/paragraph",
                raw_text=text, normalized_text=normalize_for_match(text), image_refs=refs,
            ))
            continue
        table_count += 1
        locator = f"word/{source_part}[{ordinal}]/table[{table_count}]"
        payload, refs = _table_payload(item, document, source_part, locator, image_dir, cache)
        raw_text = "\n".join(" | ".join(row) for row in payload["rows"] if any(row)).strip()
        appendix_match = _TABLE_APPENDIX_RE.search(raw_text)
        if include_in_body and appendix_match:
            section_path = _update_section_stack(section_stack, 1, appendix_match.group(1).strip())
        else:
            section_path = tuple(value for _, value in section_stack)
        blocks.append(AuditDocumentBlock(
            block_id=f"{document.document_id}-B{ordinal:04d}", document_id=document.document_id,
            ordinal=ordinal, block_type="table", section_path=section_path,
            source_locator=locator, raw_text=raw_text, normalized_text=normalize_for_match(raw_text),
            table_json=payload, image_refs=refs,
        ))
    return blocks, ordinal, paragraph_count, table_count


def parse_docx_document(document: StoredAuditDocument, docx_path: str | Path | None = None) -> ParsedAuditDocument:
    path = Path(docx_path or document.converted_path or document.original_path).resolve()
    if path.suffix.lower() != ".docx":
        raise ValueError("Word 解析器只接受 .docx 文件")
    doc = Document(str(path))
    image_dir = document.original_path.parents[1] / "images"
    cache: dict[tuple[str, str], AuditImage] = {}
    section_stack: list[tuple[int, str]] = []
    blocks, ordinal, paragraph_count, table_count = _parse_container(doc, document, "body", 0, section_stack, image_dir, cache, True)
    seen_parts: set[str] = set()
    for section_index, section in enumerate(doc.sections, start=1):
        for kind, container in (("header", section.header), ("footer", section.footer)):
            part_name = str(container.part.partname)
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            extra, ordinal, _, _ = _parse_container(container, document, f"{kind}[{section_index}]", ordinal, section_stack, image_dir, cache, False)
            blocks.extend(extra)
    body_blocks = [block for block in blocks if block.source_locator.startswith("word/body")]
    first_blocks = body_blocks[:3]
    visible_text = "".join(block.raw_text for block in first_blocks)
    cover_visual_only = bool(first_blocks and any(block.image_refs for block in first_blocks) and len(visible_text) < 20)
    return ParsedAuditDocument(document, tuple(blocks), paragraph_count, table_count, len(cache), cover_visual_only, tuple(cache.values()))
