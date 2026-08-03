import tempfile
import unittest
from pathlib import Path

from docx import Document

from kg_extract_build.audit.evidence_reader import (
    EXCEL_CELL_MAX_LENGTH, TRUNCATION_NOTICE, block_text_for_export,
    build_task_evidence_export, excel_safe_cell, resolve_group_blocks, safe_export_filename,
)
from kg_extract_build.audit.models import AuditDocumentBlock, TaskEvidenceGroup
from kg_extract_build.audit.preview import create_audit_preview


class EvidenceReaderTests(unittest.TestCase):
    def test_exports_all_tasks_as_readable_csv_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "方案.docx"
            document = Document()
            document.add_paragraph("第一章 编制依据")
            document.add_paragraph("1.1 文件依据")
            document.add_paragraph("《测试规范》")
            document.add_paragraph("第四章 施工安排")
            document.add_paragraph("4.1 施工顺序")
            document.add_paragraph("作业准备→开工条件确认→施工")
            document.save(path)
            preview = create_audit_preview("方案.docx", path.read_bytes(), storage_dir=directory)
            exported = build_task_evidence_export(preview)
            self.assertEqual(len(exported), 44)
            self.assertEqual(exported.columns.tolist()[-2:], ["人工判断", "备注"])
            self.assertTrue(exported["任务 ID"].iloc[0])

    def test_export_filename_removes_windows_illegal_characters(self):
        self.assertEqual(safe_export_filename('方案:测试.docx'), '方案_测试_任务证据.csv')

    def test_resolve_group_blocks_deduplicates_sorts_and_skips_missing_ids(self):
        first = AuditDocumentBlock("first", "doc", 2, "paragraph", (), "word/body[2]/paragraph", "第二", "第二")
        second = AuditDocumentBlock("second", "doc", 1, "paragraph", (), "word/body[1]/paragraph", "第一", "第一")
        parsed = type("Parsed", (), {"blocks": (first, second)})()
        group = TaskEvidenceGroup("group", "first", ("missing", "first", "second"), (), (), 1, "test")

        self.assertEqual([block.block_id for block in resolve_group_blocks(parsed, group)], ["second", "first"])

    def test_export_table_images_and_long_content_are_readable(self):
        table = AuditDocumentBlock(
            "table", "doc", 1, "table", (), "word/body[1]/table[1]", "fallback", "fallback",
            table_json={"rows": [["表头", ""], ["值", None]]}, image_refs=("image-1",),
        )
        self.assertEqual(block_text_for_export(table), "表头 | \n值 | \n[图片]")

        long_text = AuditDocumentBlock("long", "doc", 2, "paragraph", (), "word/body[2]/paragraph", "x" * 40_000, "x")
        self.assertTrue(len(block_text_for_export(long_text)) > EXCEL_CELL_MAX_LENGTH)
        truncated = excel_safe_cell(block_text_for_export(long_text))
        self.assertEqual(len(truncated), EXCEL_CELL_MAX_LENGTH)
        self.assertTrue(truncated.endswith(TRUNCATION_NOTICE))
