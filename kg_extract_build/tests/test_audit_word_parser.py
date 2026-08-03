import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

from kg_extract_build.audit.document_store import AuditDocumentError, store_uploaded_word
from kg_extract_build.audit.preview import create_audit_preview
from kg_extract_build.audit.word_parser import parse_docx_document


class AuditWordParserTests(unittest.TestCase):
    def _build_docx_bytes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "方案.docx"
            document = Document()
            document.add_paragraph("第一章 编制依据")
            document.add_paragraph("依据《测试规范》编制。")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "设备"
            table.cell(0, 1).text = "数量"
            table.cell(1, 0).text = "吊车"
            table.cell(1, 1).text = "1"
            document.add_paragraph("4.1 施工顺序")
            document.add_paragraph("作业准备→开工条件确认→设备安装")
            document.save(path)
            return path.read_bytes()

    def test_stores_and_parses_paragraphs_and_tables_in_body_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            stored = store_uploaded_word("测试方案.docx", self._build_docx_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertEqual(parsed.paragraph_count, 4)
            self.assertEqual(parsed.table_count, 1)
            self.assertEqual([block.block_type for block in parsed.blocks], ["heading", "paragraph", "table", "heading", "paragraph"])
            self.assertEqual(parsed.blocks[2].table_json["rows"][1], ["吊车", "1"])
            self.assertIn("第一章 编制依据", parsed.blocks[2].section_path)
            self.assertIn("4.1 施工顺序", parsed.blocks[-1].section_path)

    def test_rejects_docx_name_with_non_docx_content(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            with self.assertRaises(AuditDocumentError):
                store_uploaded_word("伪造.docx", b"not a word document", storage_dir=temp_dir)

    def test_toc_and_risk_list_do_not_expand_section_stack(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "章节.docx"
            document = Document()
            document.add_paragraph("第五章 健康、安全、环保管理\t4")
            document.add_paragraph("第五章 健康、安全、环保管理")
            document.add_paragraph("1、火灾爆炸风险控制措施")
            document.add_paragraph("1）进入井站应穿戴好防静电劳保用品")
            document.add_paragraph("5.1 风险控制要求")
            document.add_paragraph("正文内容")
            document.save(path)
            stored = store_uploaded_word("章节.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertEqual(parsed.blocks[0].block_type, "toc_entry")
            self.assertEqual(parsed.blocks[2].block_type, "paragraph")
            self.assertEqual(parsed.blocks[3].block_type, "paragraph")
            self.assertEqual(parsed.blocks[-1].section_path, ("第五章 健康、安全、环保管理", "5.1 风险控制要求"))

    def test_numbered_heading_without_space_is_recognized(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "无空格标题.docx"
            document = Document()
            document.add_paragraph("5.1工作前风险分析")
            document.add_paragraph("风险分析正文")
            document.save(path)
            stored = store_uploaded_word("无空格标题.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)
            self.assertEqual(parsed.blocks[0].block_type, "heading")
            self.assertEqual(parsed.blocks[1].section_path, ("5.1工作前风险分析",))

    def test_template_chapter_and_sentence_styled_as_heading_are_distinguished(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "标题.docx"
            document = Document()
            document.add_paragraph("施工安排")
            false_heading = document.add_paragraph("根据现场施工要求及标准准备好公示牌及施工工器具等；")
            false_heading.style = "Heading 3"
            document.add_paragraph("第五章 健康、安全、环保管理")
            document.add_paragraph("5.1 工作前风险分析")
            document.save(path)
            stored = store_uploaded_word("标题.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertEqual(parsed.blocks[0].block_type, "heading")
            self.assertEqual(parsed.blocks[1].block_type, "paragraph")
            self.assertEqual(parsed.blocks[-1].section_path, ("第五章 健康、安全、环保管理", "5.1 工作前风险分析"))

    def test_horizontal_merged_cells_are_not_repeated(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "合并单元格.docx"
            document = Document()
            table = document.add_table(rows=2, cols=3)
            table.cell(0, 0).text, table.cell(0, 1).text, table.cell(0, 2).text = "A", "B", "C"
            table.cell(1, 0).merge(table.cell(1, 2)).text = "合并说明"
            document.save(path)
            stored = store_uploaded_word("合并单元格.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertEqual(parsed.blocks[0].table_json["rows"][1], ["合并说明", "", ""])
            self.assertEqual(parsed.blocks[0].table_json["cells"][-1]["colspan"], 3)

    def test_appendix_title_in_table_starts_a_real_appendix_region(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "附录表.docx"
            document = Document()
            document.add_paragraph("附录D：设备装置评估表")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "附录E：风险作业方案控制措施检查确认表"
            table.cell(0, 1).text = "控制措施"
            document.save(path)
            stored = store_uploaded_word("附录表.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertTrue(parsed.blocks[-1].section_path[0].startswith("附录E"))

    def test_toc_entries_inside_content_controls_keep_document_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "目录控件.docx"
            document = Document()
            toc = document.add_paragraph("第一章 编制依据\t1")
            following = document.add_paragraph("第一章 编制依据")
            content = OxmlElement("w:sdtContent")
            body = toc._p.getparent()
            position = list(body).index(toc._p)
            body.remove(toc._p)
            content.append(toc._p)
            control = OxmlElement("w:sdt")
            control.append(OxmlElement("w:sdtPr"))
            control.append(content)
            body.insert(position, control)
            document.save(path)
            stored = store_uploaded_word("目录控件.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            self.assertEqual([block.block_type for block in parsed.blocks], ["toc_entry", "heading"])
            self.assertEqual(parsed.blocks[1].raw_text, following.text)

    def test_vertical_merged_cells_do_not_repeat_the_origin_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "纵向合并.docx"
            document = Document()
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "合并列"
            table.cell(1, 0).text = "应清空"
            table.cell(0, 0).merge(table.cell(1, 0))
            table.cell(0, 1).text, table.cell(1, 1).text = "A", "B"
            document.save(path)
            stored = store_uploaded_word("纵向合并.docx", path.read_bytes(), storage_dir=temp_dir)
            parsed = parse_docx_document(stored)

            rows = parsed.blocks[0].table_json["rows"]
            self.assertIn("合并列", rows[0][0])
            self.assertEqual(rows[1], ["", "B"])

    def test_preview_instantiates_location_records_for_all_published_tasks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            preview = create_audit_preview(
                "测试方案.docx", self._build_docx_bytes(), storage_dir=temp_dir
            )

            self.assertEqual(len(preview.task_library.tasks), 44)
            self.assertEqual(len(preview.locations), 44)
            self.assertNotEqual(preview.locations["ARR-002"].status, "not_located")


if __name__ == "__main__":
    unittest.main()
