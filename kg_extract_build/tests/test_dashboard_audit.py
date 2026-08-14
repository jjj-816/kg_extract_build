import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from streamlit.testing.v1 import AppTest
from kg_extract_build import dashboard_audit


class AuditDashboardTests(unittest.TestCase):
    def test_provider_widget_defaults_are_initialized_only_when_absent(self):
        state = {"audit_context_provider_base_url": "https://user.example/v1"}
        original_state = dashboard_audit.st.session_state
        try:
            dashboard_audit.st.session_state = state
            dashboard_audit._initialize_audit_provider_fields("zhipu", "zhipu")
        finally:
            dashboard_audit.st.session_state = original_state

        self.assertEqual(state["audit_context_provider_base_url"], "https://user.example/v1")
        self.assertTrue(state["audit_context_provider_model"])

    @staticmethod
    def _docx_bytes() -> bytes:
        document = Document()
        document.add_heading("施工方案", level=1)
        document.add_paragraph("本工程涉及动火作业，作业前应完成安全交底和风险确认。")
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def test_audit_page_is_available_without_mysql_connection(self):
        dashboard = Path(__file__).resolve().parents[1] / "dashboard.py"
        app = AppTest.from_file(str(dashboard)).run(timeout=15)
        app.sidebar.radio[0].set_value("施工方案审核").run(timeout=15)

        self.assertEqual(list(app.exception), [])
        self.assertTrue(any(title.value == "施工方案审核" for title in app.title))
        self.assertTrue(any(item.label == "输入已创建的 run_id" for item in app.text_input))
        self.assertTrue(any(item.label == "审核 LLM provider" for item in app.selectbox))
        self.assertTrue(any(item.label == "审核 provider Base URL" for item in app.text_input))

    def test_audit_page_upload_regression_without_mysql(self):
        dashboard = Path(__file__).resolve().parents[1] / "dashboard.py"
        app = AppTest.from_file(str(dashboard)).run(timeout=15)
        app.sidebar.radio[0].set_value("施工方案审核").run(timeout=15)
        app.file_uploader[0].set_value(
            (
                "acceptance.docx",
                self._docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        )
        app = app.run(timeout=30)

        self.assertEqual(list(app.exception), [])
        self.assertEqual(
            {tab.label for tab in app.tabs},
            {"证据块预览", "任务证据阅读器", "审核上下文"},
        )
        labels = {item.label for item in app.text_input}
        self.assertIn("方案声明规范（用逗号分隔）", labels)
        self.assertIn("作业类型必备补充规范（用逗号分隔）", labels)


if __name__ == "__main__":
    unittest.main()
