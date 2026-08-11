import os
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from streamlit.testing.v1 import AppTest


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("施工方案", level=1)
    document.add_paragraph("本工程按施工方案组织实施，作业前应完成安全交底。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class LiveAuditDashboardTests(unittest.TestCase):
    @unittest.skipUnless(os.getenv("AUDIT_LIVE_UI") == "1", "set AUDIT_LIVE_UI=1 to run against live services")
    def test_live_audit_page_reaches_stage_two(self):
        dashboard = Path(__file__).resolve().parents[1] / "dashboard.py"
        app = AppTest.from_file(str(dashboard)).run(timeout=20)
        app.sidebar.radio[0].set_value("施工方案审核").run(timeout=20)
        app.file_uploader[0].set_value(
            (
                "live-acceptance.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        )
        app = app.run(timeout=30)

        for item in app.text_input:
            if "run_id" in item.label or "API Key" in item.label:
                continue
            if "Base URL" in item.label:
                continue
            if "发布版 ID" in item.label:
                item.set_value(os.environ["KG_AUDIT_NORMATIVE_RELEASE_ID"])
                continue
            item.set_value("qwen3:0.6b")
        for item in app.text_area:
            item.set_value("验证审核页面生产链路")
        if app.selectbox:
            app.selectbox[0].set_value("ollama")
        for item in app.checkbox:
            item.set_value(True)
        app = app.run(timeout=30)
        launch = next(item for item in app.button if "确认审核上下文并创建审核运行" in item.label)
        app = launch.click().run(timeout=180)

        self.assertEqual(list(app.exception), [])
        messages = [item.value for item in (*app.error, *app.warning, *app.info, *app.success)]
        self.assertTrue(any("阶段 2" in value for value in messages), "\\n".join(messages))

        report = app.session_state["audit_stage2_report"]
        self.assertEqual(len(report["tasks"]), 44)
        self.assertTrue(app.session_state["audit_run_id"])


if __name__ == "__main__":
    unittest.main()
